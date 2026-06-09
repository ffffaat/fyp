import os
import re
from dataclasses import dataclass, field
from typing import Dict, List, Tuple

from pypdf import PdfReader
from openpyxl import load_workbook
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

try:
    from sentence_transformers import SentenceTransformer
    _sbert = SentenceTransformer("all-MiniLM-L6-v2")
except Exception:
    _sbert = None

# =========================================================
# TEXT UTILITIES
# =========================================================
# Cleans raw extracted text by removing links, emails, symbols, and extra spacing.
def clean_text(text: str) -> str:
    if not text:
        return ""
    text = re.sub(r"https?://\S+|www\.\S+", " ", text)
    text = re.sub(r"\b\S+@\S+\b", " ", text)
    text = re.sub(r"[•●■◆▶►▪]", " ", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()

# Converts text into lowercase normalized form for easier matching.
def normalize_text(text: str) -> str:
    text = text.lower()
    text = re.sub(r"[^a-z0-9\s]", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text

# Cleans and normalizes a single line of text.
def normalize_line(line: str) -> str:
    return normalize_text(clean_text(line))

# Extracts readable text from all pages of a PDF file.
def extract_pdf_text(pdf_path: str) -> str:
    reader = PdfReader(pdf_path, strict=False)
    page_texts = []
    for page in reader.pages:
        t = page.extract_text() or ""
        if t.strip():
            page_texts.append(t)
    return "\n\n".join(page_texts)

# =========================================================
# NOISE REMOVAL
# =========================================================
HEADER_FOOTER_PATTERNS = [
    r"cds505.*data visuali[sz]ation.*visual analytics",
    r"academic session",
    r"school of computer sciences",
    r"pusat pengajian sains komputer",
    r"group project",
    r"page \d+ of \d+",
    r"submission deadline",
    r"course code and title",
    r"lecturer.?s name",
    r"matric no",
    r"date issued",
    r"prepared for",
    r"prepared by",
    r"group leader",
    r"members",
]

TOC_HINTS = [
    "table of contents", "table of content", "list of figures", "list of tables", "contents"
]

# Detects whether a line is likely part of the table of contents
def is_toc_line(line: str) -> bool:
    low = line.lower()
    if any(h in low for h in TOC_HINTS):
        return True
    dotted = len(re.findall(r"\.{4,}", line))
    sec_like = len(re.findall(r"\b\d+(\.\d+)*\b", line))
    if dotted >= 1 and sec_like >= 1:
        return True
    if re.search(r"\.{4,}\s*\d+\s*$", line):
        return True
    return False

# Detects noisy lines such as headers, footers, page numbers, and repeated labels.
def is_noise_line(line: str) -> bool:
    low = normalize_line(line)
    if not low:
        return True
    if len(low.split()) <= 2 and re.fullmatch(r"[ivxlcdm\d]+", low):
        return True
    if is_toc_line(line):
        return True
    for pat in HEADER_FOOTER_PATTERNS:
        if re.search(pat, low, flags=re.I):
            return True
    return False

# Removes unnecessary header, footer, and repeated noise from extracted text.
def remove_noise(text: str) -> str:
    text = text.replace("\r", "\n")
    raw_lines = [ln.strip() for ln in text.split("\n")]
    lines = []
    for ln in raw_lines:
        if is_noise_line(ln):
            continue
        lines.append(ln)

    freq = {}
    for ln in lines:
        key = normalize_line(ln)
        if key:
            freq[key] = freq.get(key, 0) + 1

    filtered = []
    for ln in lines:
        key = normalize_line(ln)
        if freq.get(key, 0) >= 3 and len(key.split()) <= 12:
            continue
        filtered.append(ln)

    text = "\n".join(filtered)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return clean_text(text)

# Removes front matter and starts the text from the main report content.
def strip_front_matter(text: str) -> str:
    lower = text.lower()
    patterns = [
        r"\babstract\b",
        r"\b1\.0\s+introduction\b",
        r"\b1\s*\.\s*introduction\b",
        r"\b1\s+introduction\b",
        r"\bintroduction and background\b",
        r"\bbackground\b",
    ]
    starts = []
    for pat in patterns:
        m = re.search(pat, lower, flags=re.I)
        if m:
            starts.append(m.start())
    return text[min(starts):] if starts else text


# =========================================================
# RUBRIC PARSING
# =========================================================
# Reads rubric criteria, weight, and scale from the uploaded Excel rubric file.
def parse_rubric_xlsx(rubric_path: str) -> List[Dict]:
    if not rubric_path.lower().endswith((".xlsx", ".xlsm", ".xltx", ".xltm")):
        raise ValueError("Rubric file must be Excel format.")

    wb = load_workbook(rubric_path, data_only=True)
    ws = wb.active
    criteria = []

    for row in range(1, ws.max_row + 1):
        criterion = ws[f"B{row}"].value
        weight = ws[f"G{row}"].value
        max_scale = ws[f"H{row}"].value

        if criterion and isinstance(criterion, str):
            crit = criterion.strip()
            if crit.lower() in {"criteria / scale", "criteria", "criterion", "scale", "individual"}:
                continue
            try:
                weight_val = float(weight)
                max_scale_val = float(max_scale)
            except Exception:
                continue
            
            criteria.append({
                "criterion_name": crit,
                "weight": weight_val,
                "max_scale": max_scale_val,
            })

    if not criteria:
        raise ValueError("No rubric criteria found.")
    
    return criteria

# =========================================================
# SEMANTIC SCORING
# =========================================================
# Calculates semantic similarity between two text passages. Uses SBERT if available, otherwise falls back to TF-IDF cosine similarity.
def semantic_similarity(a: str, b: str) -> float:
    if not a.strip() or not b.strip():
        return 0.0

    if _sbert is not None:
        emb = _sbert.encode([a[:3000], b[:3000]])
        return max(0.0, float(cosine_similarity([emb[0]], [emb[1]])[0][0]))

    vect = TfidfVectorizer(stop_words="english", ngram_range=(1, 2), max_features=2000)
    X = vect.fit_transform([a, b])
    return max(0.0, float(cosine_similarity(X[0:1], X[1:2])[0][0]))

# Calculates how many required concept terms appear in the student text.
def keyword_coverage(student_text: str, concept_terms: List[str]) -> float:
    if not concept_terms:
        return 0.0
    st = normalize_text(student_text)
    hits = 0
    for term in concept_terms:
        term_n = normalize_text(term)
        if term_n and term_n in st:
            hits += 1
    return hits / len(concept_terms)

# Counts the number of exact term matches found in the text.
def exact_term_hits(text: str, terms: List[str]) -> int:
    t = normalize_text(text)
    return sum(1 for term in terms if normalize_text(term) and normalize_text(term) in t)

# =========================================================
# SECTION DETECTION / PROFILES
# =========================================================
@dataclass
class CriterionProfile:
    name: str
    concept_groups: List[List[str]] = field(default_factory=list)
    optional_terms: List[str] = field(default_factory=list)
    section_aliases: List[str] = field(default_factory=list)
    score_mode: str = "semantic"

# =========================================================
# CRITERION NAME NORMALIZATION
# =========================================================
CRITERION_NAME_ALIASES: Dict[str, List[str]] = {
    "Abstract, Intro & Background": [
        "abstract, intro & background", "abstract", "introduction",
        "background", "introduction and background", "abstract introduction and background",
        "introduction & background", "problem statement", "project background",
        "big picture", "real world problem", "overview and background",
        "introduction and overview", "context and background", "motivation and background",
        "problem and motivation", "dataset description and preparation", "background and dataset description",
        "introduction background and dataset", "project objectives", "project objective",
    ],

    "Data & Task Abstraction": [
        "data & task abstraction", "data abstraction", "task abstraction",
        "data and task abstraction", "data/task abstraction", "dataset abstraction",
        "data analysis task", "task analysis", "analytical tasks",
        "what and why", "munzner abstraction", "data types and tasks", 
        "dataset types and task abstraction", "data abstraction and task analysis", "what why abstraction",
        "data description and task abstraction", "dataset structure and tasks", "data understanding and task analysis",
    ],

    "Methods & Tools for Data Storytelling": [
        "methods & tools for data storytelling", "methods and tools for data storytelling", "methods & tools",
        "methods and tools", "tools and methods", "methodology and tools",
        "design choice and visual encoding", "visual encoding and design choice", "design justification",
        "visual design justification", "implementation and tools", "tools used",
        "software and methods", "storytelling approach", "data storytelling approach",
        "design and implementation", "methods tools and storytelling", "design choice & encoding",
        "design choice and encoding", "visual mapping and tools",
    ],

    "Methods & Tools": [
        "methods & tools", "methods and tools", "tools and methods",
        "methodology", "methods", "tools", "implementation",
        "implementation details", "software used", "workflow", "analysis workflow",
        "preprocessing and tools", "data preparation and tools",
    ],

    "Related Work": [
        "related work", "literature review", "review of literature", "previous studies",
        "prior work", "related studies", "research background", "research context", "comparison with existing work",
        "existing studies", "state of the art", "related literature",
    ],

    "Discussion and Conclusion": [
        "discussion and conclusion", "discussion", "conclusion", "summary and conclusion",
        "discussion, limitation and conclusion", "discussion and findings", "findings and conclusion",
        "results and discussion", "discussion results and conclusion", "critical reflection",
        "reflection and conclusion", "limitations and future work", "conclusion and future work",
        "findings limitations and future work", "final discussion", "discussion and limitations", "reflection",
    ],

    "References": [
        "references", "reference", "reference list", "bibliography",
        "works cited", "citations", "sources", "list of references",
    ],

    "Vis Tool Project File & Datasets": [
        "vis tool project file & datasets", "vis tool project file and datasets", "project file and datasets",
        "visualization tool project file and datasets", "visualisation tool project file and datasets", "submission package",
        "dataset and project file", "project files", "dataset files", "tool and dataset files",
        "data source and project file", "files and datasets", "appendix project files",
    ],

    "Visual Analytics": [
        "visual analytics", "interactive visual analytics", "visual analysis", "dashboard design",
        "visualization design", "visualisation design", "dashboard and interaction", "interactive dashboard design",
        "interaction design", "visual encoding and analytics", "analytics strategy", "exploratory visual analysis",
        "visual design", "dashboard analysis", "interactive dashboard",
    ],

    "Future Improvement": [
        "future improvement", "future improvements", "future work", "recommendations",
        "recommendation", "future direction", "next steps", "improvements and future work",
    ],
}

# Normalizes different rubric criterion names into standard internal names.
def normalize_criterion_name(raw_name: str) -> str:
    if not raw_name:
        return raw_name

    norm = re.sub(r"\s+", " ", raw_name.strip().lower())

    for canonical_name, aliases in CRITERION_NAME_ALIASES.items():
        alias_norms = [re.sub(r"\s+", " ", a.strip().lower()) for a in aliases]
        if norm in alias_norms:
            return canonical_name

    for canonical_name, aliases in CRITERION_NAME_ALIASES.items():
        alias_norms = [re.sub(r"\s+", " ", a.strip().lower()) for a in aliases]
        for alias in alias_norms:
            if alias in norm or norm in alias:
                return canonical_name

    return raw_name

DEFAULT_PROFILES: Dict[str, CriterionProfile] = {
    "Abstract, Intro & Background": CriterionProfile(
        name="Abstract, Intro & Background",
        concept_groups=[
            ["problem", "issue", "challenge"],
            ["objective", "aim", "goal", "purpose"],
            ["dataset", "data source", "corpus", "sample"],
            ["motivation", "importance", "significance", "relevance"],
        ],
        optional_terms=["background", "introduction", "scope", "context", "target audience"],
        section_aliases=["abstract", "introduction", "background", "target audience", "big picture", "real world problem", "dataset description"],
        score_mode="semantic",
    ),
    "Data & Task Abstraction": CriterionProfile(
        name="Data & Task Abstraction",
        concept_groups=[
            ["data type", "attribute", "feature", "variable"],
            ["task", "analysis task", "comparison", "trend", "distribution"],
            ["derived", "transformed", "preprocessing", "cleaning"],
            ["user", "audience", "stakeholder"],
        ],
        optional_terms=["what", "why", "how", "abstraction", "dataset type", "task abstraction"],
        section_aliases=["data abstraction", "task abstraction", "dataset types", "data types", "what", "why"],
        score_mode="semantic",
    ),
    "Methods & Tools for Data Storytelling": CriterionProfile(
        name="Methods & Tools for Data Storytelling",
        concept_groups=[
            ["method", "approach", "technique"],
            ["tool", "software", "tableau", "python", "power bi"],
            ["visualization", "dashboard", "chart", "story"],
            ["interaction", "filter", "parameter", "action"],
        ],
        optional_terms=["implementation", "design", "justification", "storytelling"],
        section_aliases=["methods", "tools", "design justification", "visual encoding", "implementation", "tools used", "design choice"],
        score_mode="tool_dataset",
    ),
    "Methods & Tools": CriterionProfile(
        name="Methods & Tools",
        concept_groups=[
            ["method", "approach", "technique"],
            ["tool", "software", "tableau", "python", "power bi"],
        ],
        optional_terms=["implementation", "workflow", "tool", "software"],
        section_aliases=["methods", "tools", "implementation", "data cleaning", "tools used"],
        score_mode="tool_dataset",
    ),
    "Related Work": CriterionProfile(
        name="Related Work",
        concept_groups=[
            ["literature", "related work", "previous study", "prior work"],
            ["comparison", "similar", "difference", "gap"],
            ["citation", "reference", "source"],
        ],
        optional_terms=["author", "year", "journal", "paper"],
        section_aliases=["related work", "literature review", "previous studies", "research background"],
        score_mode="semantic",
    ),
    "Discussion and Conclusion": CriterionProfile(
        name="Discussion and Conclusion",
        concept_groups=[
            ["finding", "result", "outcome"],
            ["discussion", "interpretation", "insight"],
            ["limitation", "constraint", "weakness"],
            ["conclusion", "future work", "recommendation"],
        ],
        optional_terms=["summary", "overall", "reflection"],
        section_aliases=["discussion", "conclusion", "reflection", "future work", "what worked", "limitations", "future improvements", "critical reflection"],
        score_mode="semantic",
    ),
    "References": CriterionProfile(
        name="References",
        concept_groups=[["reference", "citation", "bibliography"]],
        optional_terms=["doi", "journal", "conference"],
        section_aliases=["references", "bibliography", "works cited"],
        score_mode="references",
    ),
    "Vis Tool Project File & Datasets": CriterionProfile(
        name="Vis Tool Project File & Datasets",
        concept_groups=[
            ["dataset", "data source", "csv", "excel"],
            ["tableau", "tool", "software", "dashboard"],
            ["cleaning", "preprocessing", "preparation"],
        ],
        optional_terms=["file", "dataset", "visualization tool", "tableau"],
        section_aliases=["data cleaning", "data source", "dataset", "design justification", "tools and technologies", "submission package", "project file"],
        score_mode="tool_dataset",
    ),
    "Visual Analytics": CriterionProfile(
        name="Visual Analytics",
        concept_groups=[
            ["visual analytics", "interactive", "dashboard"],
            ["trend", "comparison", "outlier", "relationship"],
            ["encoding", "visual", "analysis"],
        ],
        optional_terms=["interaction", "analytics", "visualisation", "visualization"],
        section_aliases=["visual analytics", "design justification", "visual encoding", "interaction design", "visual encoding and analytics", "analytics strategy"],
        score_mode="semantic",
    ),
    "Future Improvement": CriterionProfile(
        name="Future Improvement",
        concept_groups=[
            ["future", "future work", "future improvement", "enhancement"],
            ["recommendation", "improve", "extension"],
            ["limitation", "constraint", "weakness"],
        ],
        optional_terms=["next step", "improvement", "further study", "future direction"],
        section_aliases=["future improvement", "future improvements", "future work", "recommendation", "next steps"],
        score_mode="semantic",
    ),
}

# Retrieves the scoring profile for a rubric criterion.
def resolve_profile(criterion_name: str) -> CriterionProfile:
    criterion_name = normalize_criterion_name(criterion_name)

    if criterion_name in DEFAULT_PROFILES:
        return DEFAULT_PROFILES[criterion_name]

    generic_terms = [w for w in re.split(r"\W+", criterion_name.lower()) if len(w) > 3]
    return CriterionProfile(
        name=criterion_name,
        concept_groups=[[term] for term in generic_terms[:5]],
        optional_terms=generic_terms[5:10],
        section_aliases=generic_terms[:3],
        score_mode="semantic",
    )

# Detects whether a line is likely to be a section heading.
def is_heading_line(line: str) -> bool:
    s = clean_text(line)
    if not s:
        return False
    if len(s.split()) > 10:
        return False
    if re.match(r"^\d+(\.\d+)*\s+[A-Za-z]", s):
        return True
    low = s.lower()
    heading_words = [
        "abstract", "introduction", "background", "data abstraction", "task abstraction",
        "methods", "tools", "related work", "discussion", "conclusion", "references",
        "design justification", "visual analytics", "validation", "target audience",
        "data cleaning", "dataset types", "data types", "writing style", "what worked",
        "future improvements", "critical reflection", "tools used", "project objectives"
    ]
    return any(hw == low or low.startswith(hw) for hw in heading_words)

# Converts a heading line into a normalized section title.
def canonical_heading(line: str) -> str:
    s = clean_text(line)
    s = re.sub(r"^\d+(\.\d+)*\s*", "", s)
    return normalize_text(s)

# Extracts report sections based on detected headings.
def extract_sections(text: str) -> Dict[str, str]:
    text = remove_noise(text)
    text = strip_front_matter(text)
    lines = [ln.strip() for ln in text.split("\n") if ln.strip()]

    sections: Dict[str, List[str]] = {"full_document": []}
    current = "full_document"
    for ln in lines:
        if is_heading_line(ln):
            current = canonical_heading(ln)
            sections.setdefault(current, [])
            continue
        sections.setdefault(current, []).append(ln)
    return {k: clean_text("\n".join(v)) for k, v in sections.items() if clean_text("\n".join(v))}

# Selects the most relevant report section for a criterion.
def best_section_text(sections: Dict[str, str], profile: CriterionProfile, fallback_full_text: str) -> Tuple[str, float, List[str]]:
    if not sections:
        return fallback_full_text, 0.2, []

    candidates = []
    alias_terms = list(profile.section_aliases or []) + list(profile.optional_terms) + [profile.name]
    body_probe_terms = [t for grp in profile.concept_groups for t in grp[:2]] + profile.optional_terms[:6]

    for title, body in sections.items():
        if title == "full_document" or not body.strip():
            continue
        norm_title = normalize_text(title)
        title_hits = sum(1 for alias in alias_terms if normalize_text(alias) and normalize_text(alias) in norm_title)
        body_score = keyword_coverage(body, body_probe_terms)
        combined = (0.50 * min(title_hits, 3) / 3.0) + (0.50 * body_score)
        candidates.append((title, body, combined))

    if not candidates:
        return fallback_full_text, 0.2, []

    candidates.sort(key=lambda x: x[2], reverse=True)
    top = candidates[:2]
    best_conf = top[0][2]

    if best_conf < 0.55:
        return fallback_full_text, round(best_conf, 3), ["full document fallback"]

    merged = "\n\n".join([t[1] for t in top])
    titles = [t[0] for t in top]
    return merged if merged.strip() else fallback_full_text, round(best_conf, 3), titles

# =========================================================
# CHUNKING / RETRIEVAL
# =========================================================
# Splits long text into smaller chunks for retrieval and scoring.
def split_into_chunks(text: str, max_words: int = 90) -> List[str]:
    text = text.replace("\r", "\n")
    paras = [clean_text(p) for p in re.split(r"\n\s*\n+", text) if clean_text(p)]

    if len(paras) <= 1:
        sentences = [s.strip() for s in re.split(r"(?<=[.!?])\s+", text) if s.strip()]
        paras = []
        buf = []
        for s in sentences:
            buf.append(s)
            if len(" ".join(buf).split()) >= max_words:
                paras.append(" ".join(buf))
                buf = []
        if buf:
            paras.append(" ".join(buf))

    final_chunks = []
    for p in paras:
        words = p.split()
        if len(words) <= max_words:
            final_chunks.append(p)
        else:
            step = max_words
            for i in range(0, len(words), step):
                final_chunks.append(" ".join(words[i:i + max_words]))
    return final_chunks

# Applies penalty to chunks that contain mixed or irrelevant sections.
def mixed_section_penalty(chunk: str) -> float:
    low = normalize_text(chunk)
    foreign_heads = [
        "references", "bibliography", "table of contents", "list of figures", "prepared for", "prepared by"
    ]
    hits = sum(1 for h in foreign_heads if h in low)
    if hits >= 2:
        return 0.65
    if hits == 1:
        return 0.85
    return 1.0

# Builds a search query from the criterion profile.
def build_query(profile: CriterionProfile, criterion_name: str) -> Tuple[str, List[str]]:
    terms = [criterion_name] + [term for grp in profile.concept_groups for term in grp[:2]] + profile.optional_terms + profile.section_aliases
    terms = [t for t in terms if t]
    return " ".join(terms), terms

# Retrieves the most relevant chunks for the selected criterion.
def retrieve_top_chunks(query: str, query_terms: List[str], chunks: List[str], top_k: int = 3) -> List[str]:
    if not chunks:
        return []
    scored = []
    for chunk in chunks:
        sem = semantic_similarity(query, chunk)
        lex = keyword_coverage(chunk, query_terms)
        bonus = min(1.0, exact_term_hits(chunk, query_terms) / max(1, len(query_terms) // 4 or 1))
        score = (0.55 * sem) + (0.30 * lex) + (0.15 * bonus)
        score *= mixed_section_penalty(chunk)
        scored.append((chunk, score))
    scored.sort(key=lambda x: x[1], reverse=True)
    top = [chunk for chunk, _ in scored[:top_k] if chunk.strip()]
    return top or chunks[:top_k]

# =========================================================
# FEATURE SCORING
# =========================================================
# Measures how many concept groups are covered in the student text.
def concept_group_coverage(text: str, concept_groups: List[List[str]]) -> Tuple[float, List[Dict]]:
    text_n = normalize_text(text)
    if not concept_groups:
        return 0.0, []

    matched_groups = 0
    details = []
    for group in concept_groups:
        group_hit = False
        matched_term = None
        for term in group:
            if normalize_text(term) in text_n:
                group_hit = True
                matched_term = term
                break
        if group_hit:
            matched_groups += 1
        details.append({"group": group, "matched": group_hit, "matched_term": matched_term})
    return matched_groups / len(concept_groups), details

# Scores the breadth of explanation based on text length.
def breadth_score(student_text: str) -> float:
    n = len(student_text.split())
    if n >= 220:
        return 1.0
    if n >= 160:
        return 0.85
    if n >= 110:
        return 0.70
    if n >= 70:
        return 0.50
    return 0.25

# Scores whether the text contains evidence or analytical indicators.
def evidence_score(student_text: str) -> float:
    text_n = normalize_text(student_text)
    indicators = [
        "for example", "such as", "because", "therefore", "this shows",
        "result", "analysis", "dataset", "figure", "table", "based on"
    ]
    hits = sum(1 for s in indicators if s in text_n)
    return min(1.0, hits / 5)

# =========================================================
# REFERENCE QUALITY CHECKER
# =========================================================
VALID_SOURCE_HINTS = [
    ".gov", ".edu", ".ac.", "doi.org", "sciencedirect", "springer", "ieee",
    "acm", "nature", "wiley", "tandfonline", "elsevier", "mdpi", "arxiv", "frontiers", "plos"
]
WEAK_SOURCE_HINTS = [
    "wikipedia", "studocu", "coursehero", "ukessays", "123dok", "blogspot",
    "wordpress", "medium.com", "scribd", "slideshare"
]
REFERENCE_SECTION_PATTERNS = [r"\breferences\b", r"\bbibliography\b", r"\bworks cited\b"]
IN_TEXT_CITATION_PATTERNS = [
    r"\([A-Z][A-Za-z\-]+,\s*\d{4}\)",
    r"\[[0-9]{1,3}\]",
    r"\([A-Z][A-Za-z\-]+\s+et\s+al\.,\s*\d{4}\)",
    r"[A-Z][A-Za-z\-]+\s+et\s+al\.\s*\(\d{4}\)",
    r"[A-Z][A-Za-z\-]+\s*\(\d{4}\)",
]

# Extracts the reference section from the full student report.
def extract_reference_section(full_text: str) -> str:
    lower = full_text.lower()
    positions = []
    for pat in REFERENCE_SECTION_PATTERNS:
        m = re.search(pat, lower, flags=re.I)
        if m:
            positions.append(m.start())
    return full_text[min(positions):] if positions else ""

# Detects whether a line is likely the start of a reference entry.
def is_likely_reference_start(line: str) -> bool:
    line = line.strip()

    if not line:
        return False

    if re.match(r"^(\[\d+\]|\d+\.)\s+", line):
        return True

    if re.match(r"^[A-Z][A-Za-z\-']+,\s+[A-Z]\.", line):
        return True

    if re.match(r"^[A-Z][A-Za-z\-']+(?:\s+[A-Z][A-Za-z\-']+){1,3},", line):
        return True

    if re.match(r"^[A-Z][A-Za-z\-']+,\s+[A-Z]\..*&", line):
        return True

    return False

# Checks whether the current reference buffer contains a publication year.
def buffer_has_reference_year(buffer_lines: List[str]) -> bool:
    joined = " ".join(buffer_lines)
    return bool(re.search(r"\(\d{4}\)|\b(19|20)\d{2}\b", joined))

# Splits the reference section into individual reference entries.
def split_reference_entries(reference_section: str) -> List[str]:
    if not reference_section:
        return []

    ref_section = remove_noise(reference_section)

    ref_section = re.sub(
        r"^\s*(references|bibliography|works cited)\s*",
        "",
        ref_section,
        flags=re.I
    ).strip()

    lines = [ln.strip() for ln in ref_section.split("\n") if ln.strip()]

    entries = []
    buffer = []

    for line in lines:
        clean_line = clean_text(line)

        if not clean_line:
            continue

        if normalize_text(clean_line) in {"references", "bibliography", "works cited"}:
            continue

        starts_new = is_likely_reference_start(clean_line)

        if starts_new and buffer and buffer_has_reference_year(buffer):
            entries.append(clean_text(" ".join(buffer)))
            buffer = [clean_line]
        else:
            buffer.append(clean_line)

    if buffer:
        entries.append(clean_text(" ".join(buffer)))

    final_entries = []
    seen = set()

    for entry in entries:
        entry = clean_text(entry)

        if not re.search(r"\(\d{4}\)|\b(19|20)\d{2}\b", entry):
            continue

        if len(entry.split()) < 8:
            continue

        key = normalize_text(entry[:160])
        if key not in seen:
            seen.add(key)
            final_entries.append(entry)

    return final_entries

# Scores one reference entry based on author, year, title, DOI/link, and source quality.
def score_reference_entry(entry: str, current_year: int = 2026) -> Dict:
    e_low = entry.lower()
    score = 0.0
    reasons = []

    has_author = bool(re.search(r"^[A-Z][A-Za-z\-']+", entry))
    has_year = bool(re.search(r"\b(19|20)\d{2}\b", entry))
    has_title = len(entry.split()) >= 6
    has_link = bool(re.search(r"https?://|www\.", entry))
    has_doi = bool(re.search(r"doi\.org|\bdoi\b", e_low))

    if has_author:
        score += 0.20; reasons.append("author-present")
    if has_year:
        score += 0.20; reasons.append("year-present")
    if has_title:
        score += 0.15; reasons.append("title-present")
    if has_link or has_doi:
        score += 0.15; reasons.append("traceable-link-or-doi")
    if any(hint in e_low for hint in VALID_SOURCE_HINTS):
        score += 0.20; reasons.append("credible-domain-or-publisher")
    if any(hint in e_low for hint in WEAK_SOURCE_HINTS):
        score -= 0.25; reasons.append("weak-source-flag")

    full_years = re.findall(r"\b((?:19|20)\d{2})\b", entry)
    if full_years:
        latest = max(int(y) for y in full_years)
        age = current_year - latest
        if age <= 5:
            score += 0.10; reasons.append("recent-source")
        elif age > 12:
            score -= 0.08; reasons.append("older-source")

    score = max(0.0, min(1.0, score))
    label = "high" if score >= 0.75 else "moderate" if score >= 0.45 else "low"
    return {"entry": entry, "score": round(score, 3), "label": label, "reasons": reasons}

# Extracts short text snippets containing in-text citations.
def extract_citation_snippets(full_text: str, window: int = 60) -> List[str]:
    snippets = []

    for pat in IN_TEXT_CITATION_PATTERNS:
        for m in re.finditer(pat, full_text):
            start = max(0, m.start() - window)
            end = min(len(full_text), m.end() + window)

            snippet = full_text[start:end].strip().replace("\n", " ")
            citation = m.group(0)
            highlighted = snippet.replace(
                citation,
                f"<mark>{citation}</mark>"
            )
            snippets.append(highlighted)

    seen = set()
    unique_snippets = []
    for s in snippets:
        if s not in seen:
            seen.add(s)
            unique_snippets.append(s)

    return unique_snippets[:10]

# Evaluates the overall reference quality of the student report.
def evaluate_references(full_text: str) -> Dict:
    ref_section = extract_reference_section(full_text)
    entries = split_reference_entries(ref_section)
    entry_scores = [score_reference_entry(e) for e in entries]

    citation_matches = []
    for pat in IN_TEXT_CITATION_PATTERNS:
        found = re.findall(pat, full_text)
        if found:
            citation_matches.extend(found)

    seen = set()
    unique_citations = []
    for c in citation_matches:
        if c not in seen:
            seen.add(c)
            unique_citations.append(c)

    in_text_count = len(citation_matches)

    avg_quality = sum(x["score"] for x in entry_scores) / len(entry_scores) if entry_scores else 0.0
    weak_count = sum(1 for x in entry_scores if x["label"] == "low")
    strong_count = sum(1 for x in entry_scores if x["label"] == "high")
    citation_alignment = min(1.0, in_text_count / max(1, len(entries))) if entries else 0.0
    overall = (0.55 * avg_quality) + (0.25 * citation_alignment) + (0.20 * min(1.0, strong_count / max(1, len(entries))))

    return {
        "reference_entries_found": len(entries),
        "in_text_citation_snippets": extract_citation_snippets(full_text),
        "in_text_citations_found": in_text_count,
        "in_text_citation_examples": unique_citations[:20],
        "average_reference_quality": round(avg_quality, 3),
        "citation_alignment": round(citation_alignment, 3),
        "weak_reference_count": weak_count,
        "strong_reference_count": strong_count,
        "overall_reference_score": round(overall, 3),
        "reference_details": entry_scores,
    }

# =========================================================
# CRITERION SCORING
# =========================================================
# Scores writing style based on coherence, repetition, and breadth.
def writing_style_score(text: str, max_scale: float) -> Dict:
    sentences = [s.strip() for s in re.split(r"[.!?]+", text) if s.strip()]
    words = text.split()
    unique_ratio = len(set(normalize_text(text).split())) / max(1, len(normalize_text(text).split()))
    avg_sentence_len = len(words) / max(1, len(sentences))

    coherence = 1.0 if 8 <= avg_sentence_len <= 28 else 0.7 if 5 <= avg_sentence_len <= 35 else 0.4
    repetition = 1.0 if unique_ratio >= 0.45 else 0.7 if unique_ratio >= 0.35 else 0.4
    breadth = breadth_score(text)
    raw = round(((0.40 * coherence) + (0.35 * repetition) + (0.25 * breadth)) * max_scale, 2)

    return {
        "best_reference_similarity": 0.0,
        "aggregate_similarity": 0.0,
        "concept_coverage": 0.0,
        "optional_term_coverage": round(repetition, 4),
        "breadth_score": round(breadth, 4),
        "evidence_score": round(coherence, 4),
        "raw_score": raw,
        "concept_details": [],
        "retrieval_confidence": 1.0,
    }

# Scores the reference criterion independently from the model answer.
def references_criterion_score(full_student_text: str, max_scale: float) -> Tuple[Dict, Dict]:
    ref = evaluate_references(full_student_text)
    raw = round(ref["overall_reference_score"] * max_scale, 2)
    metrics = {
        "best_reference_similarity": 0.0,
        "aggregate_similarity": 0.0,
        "concept_coverage": 1.0 if ref["reference_entries_found"] > 0 else 0.0,
        "optional_term_coverage": ref["citation_alignment"],
        "breadth_score": 1.0 if ref["reference_entries_found"] >= 8 else 0.7 if ref["reference_entries_found"] >= 4 else 0.3,
        "evidence_score": ref["average_reference_quality"],
        "raw_score": raw,
        "concept_details": [],
        "retrieval_confidence": 1.0,
    }
    return metrics, ref

# Calculates a blended criterion score using similarity, coverage, breadth, and evidence.
def blended_criterion_score(student_text: str, reference_texts: List[str], profile: CriterionProfile, max_scale: float, retrieval_confidence: float) -> Dict:
    student_text = clean_text(student_text)
    merged_reference = " ".join([clean_text(t) for t in reference_texts if clean_text(t)])

    ref_similarities = [semantic_similarity(student_text, r) for r in reference_texts if clean_text(r)]
    best_reference_similarity = max(ref_similarities) if ref_similarities else 0.0
    aggregate_similarity = semantic_similarity(student_text, merged_reference) if merged_reference else 0.0

    concept_cov, concept_details = concept_group_coverage(student_text, profile.concept_groups)
    optional_cov = keyword_coverage(student_text, profile.optional_terms)
    breadth = breadth_score(student_text)
    evidence = evidence_score(student_text)

    blended = (
        0.2 * best_reference_similarity +
        0.2 * aggregate_similarity +
        0.3 * concept_cov +
        0.1 * optional_cov +
        0.1 * breadth +
        0.1 * evidence
    )
    blended *= max(0.60, min(1.0, 0.65 + 0.35 * retrieval_confidence))
    raw_score = round(blended * max_scale, 2)

    return {
        "best_reference_similarity": round(best_reference_similarity, 4),
        "aggregate_similarity": round(aggregate_similarity, 4),
        "concept_coverage": round(concept_cov, 4),
        "optional_term_coverage": round(optional_cov, 4),
        "breadth_score": round(breadth, 4),
        "evidence_score": round(evidence, 4),
        "raw_score": raw_score,
        "concept_details": concept_details,
        "retrieval_confidence": round(retrieval_confidence, 4),
    }

# Builds short AI feedback for each rubric criterion.
def build_feedback(criterion_name: str, metrics: Dict, max_scale: float, section_titles: List[str]) -> str:
    score = metrics["raw_score"]
    if score >= max_scale * 0.85:
        level = "Strong"
    elif score >= max_scale * 0.60:
        level = "Moderate"
    else:
        level = "Weak"

    missing_groups = [
        "/".join(d["group"][:2])
        for d in metrics.get("concept_details", [])
        if not d["matched"]
    ]
    missing_text = ", ".join(missing_groups[:3]) if missing_groups else "none"
    section_text = ", ".join(section_titles) if section_titles else "full document fallback"

    return (
        f"{level} for '{criterion_name}'. "
        f"Best-reference similarity={metrics['best_reference_similarity']:.2f}, "
        f"aggregate similarity={metrics['aggregate_similarity']:.2f}, "
        f"concept coverage={metrics['concept_coverage']:.2f}, breadth={metrics['breadth_score']:.2f}. "
        f"Missing concept areas: {missing_text}. "
        f"Retrieval confidence={metrics.get('retrieval_confidence', 0):.2f}. Sections used: {section_text}."
    )


# =========================================================
# MAIN PIPELINE
# =========================================================
# Main pipeline that evaluates a student submission using rubric and reference answer files.
def evaluate_submission(student_pdf_path: str, rubric_xlsx_path: str, reference_pdf_paths: list[str]) -> dict:
    print("USING CHUNK TUNED NLP_ASSIGNMENT")
    if not os.path.exists(student_pdf_path):
        raise FileNotFoundError(f"Student submission not found: {student_pdf_path}")
    if not os.path.exists(rubric_xlsx_path):
        raise FileNotFoundError(f"Rubric file not found: {rubric_xlsx_path}")

    reference_pdf_paths = reference_pdf_paths or []
    for p in reference_pdf_paths:
        if not os.path.exists(p):
            raise FileNotFoundError(f"Reference file not found: {p}")

    student_text_raw = extract_pdf_text(student_pdf_path)
    student_text = strip_front_matter(remove_noise(student_text_raw))
    rubric_items = parse_rubric_xlsx(rubric_xlsx_path)

    reference_texts_raw = [extract_pdf_text(p) for p in reference_pdf_paths]
    reference_texts = [strip_front_matter(remove_noise(t)) for t in reference_texts_raw]

    if not student_text.strip():
        raise ValueError("Student PDF text extraction failed.")

    student_sections = extract_sections(student_text)
    reference_sections_by_doc = [extract_sections(t) for t in reference_texts]

    results = []
    total_weight = sum(item["weight"] for item in rubric_items) or 1.0
    total_weighted_score = 0.0
    references_report = evaluate_references(student_text)

    for item in rubric_items:
        criterion_name_raw = item["criterion_name"]
        criterion_name = normalize_criterion_name(criterion_name_raw)
        weight = item["weight"]
        max_scale = item["max_scale"]
        profile = resolve_profile(criterion_name)
        query, query_terms = build_query(profile, criterion_name)

        selected_student_section_text, retrieval_conf, used_student_titles = best_section_text(student_sections, profile, student_text)

        if criterion_name in {"Related Work", "Abstract, Intro & Background", "Visual Analytics"} and retrieval_conf < 0.52:
            selected_student_section_text = student_text
            used_student_titles = ["full document fallback"]

        student_section_chunks = split_into_chunks(selected_student_section_text, max_words=90)
        selected_student_chunks = retrieve_top_chunks(query, query_terms, student_section_chunks, top_k=3)
        if not selected_student_chunks:
            selected_student_chunks = [selected_student_section_text] if selected_student_section_text else []

        selected_reference_texts = []
        selected_reference_chunks = []
        used_reference_titles = []

        for ref_text, ref_sections in zip(reference_texts, reference_sections_by_doc):
            if profile.score_mode in {"references", "writing"}:
                selected_reference_chunks.append([])
                selected_reference_texts.append("")
                continue

            section_text, ref_conf, ref_titles = best_section_text(ref_sections, profile, ref_text)
            if criterion_name in {"Related Work", "Abstract, Intro & Background", "Visual Analytics"} and ref_conf < 0.52:
                section_text = ref_text
                ref_titles = ["full document fallback"]

            ref_section_chunks = split_into_chunks(section_text, max_words=90)
            ref_chunks = retrieve_top_chunks(query, query_terms, ref_section_chunks, top_k=2)
            if not ref_chunks:
                ref_chunks = [section_text] if section_text else []
            selected_reference_chunks.append(ref_chunks)
            selected_reference_texts.append(" ".join(ref_chunks))
            used_reference_titles.extend(ref_titles)

        reference_review = None
        if profile.score_mode == "references":
            metric_result, reference_review = references_criterion_score(student_text, max_scale)
            selected_reference_chunks = []
            selected_student_chunks = []
            used_student_titles = ["independent reference review"]
        elif profile.score_mode == "writing":
            metric_result = writing_style_score(student_text, max_scale)
        else:
            metric_result = blended_criterion_score(
                student_text=" ".join(selected_student_chunks),
                reference_texts=selected_reference_texts,
                profile=profile,
                max_scale=max_scale,
                retrieval_confidence=retrieval_conf,
            )

        weighted_score = round((metric_result["raw_score"] / max_scale) * weight, 2) if max_scale else 0.0
        total_weighted_score += weighted_score

        if profile.score_mode == "references" and reference_review is not None:
            feedback = (
                "Reference quality scored independently from the model answer. "
                f"Entries={reference_review['reference_entries_found']}, "
                f"in-text citations={reference_review['in_text_citations_found']}, "
                f"average quality={reference_review['average_reference_quality']:.2f}, "
                f"citation alignment={reference_review['citation_alignment']:.2f}, "
                f"weak sources={reference_review['weak_reference_count']}, "
                f"strong sources={reference_review['strong_reference_count']}."
            )
        else:
            feedback = build_feedback(criterion_name, metric_result, max_scale, used_student_titles)

        results.append({
            "criterion_name": criterion_name,              
            "criterion_name_raw": criterion_name_raw,      
            "weight": weight,
            "max_scale": max_scale,
            "raw_score": metric_result["raw_score"],
            "weighted_score": weighted_score,
            "feedback": feedback,
            "student_chunks": selected_student_chunks,
            "reference_chunks": selected_reference_chunks,
            "student_section_titles": used_student_titles,
            "reference_section_titles": used_reference_titles,
            "is_reference_independent": profile.score_mode == "references",
            "reference_review": reference_review,
            **metric_result,
        })

    final_content_score = round((total_weighted_score / total_weight) * 100, 2)
    reference_quality_score = round(references_report["overall_reference_score"] * 100, 2)
    final_score = round((0.9 * final_content_score) + (0.1 * reference_quality_score), 2)

    return {
        "final_score": final_score,
        "final_content_score": final_content_score,
        "reference_quality_score": reference_quality_score,
        "criteria_results": results,
        "reference_report": references_report,
        "student_text": student_text,
        "reference_text_count": len(reference_texts),
    }
